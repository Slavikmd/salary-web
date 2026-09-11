#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import json
import sqlite3
import hashlib
import calendar
from datetime import datetime, timedelta
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO

app = Flask(__name__)
app.secret_key = 'your-secret-key-change-this-2026'

# === БАЗА ДАННЫХ ===

DB_PATH = os.path.join(os.path.dirname(__file__), 'salary.db')

def init_db():
    """Инициализация базы данных"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Таблица пользователей
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # Таблица записей
    c.execute('''
        CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            date TEXT NOT NULL,
            work TEXT,
            hours REAL,
            salary REAL,
            salary_day INTEGER DEFAULT 0,
            comment TEXT,
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(user_id, date)
        )
    ''')
    
    conn.commit()
    conn.close()

def hash_password(password):
    """Хеширование пароля"""
    return hashlib.sha256(password.encode()).hexdigest()

def get_db():
    """Подключение к базе данных"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

# === ДЕКОРАТОР АВТОРИЗАЦИИ ===

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

# === МАРШРУТЫ ===

@app.route('/')
@login_required
def index():
    return render_template('index.html', username=session.get('username'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        
        if not username or not password:
            return render_template('login.html', error='Заполните все поля')
        
        conn = get_db()
        user = conn.execute(
            'SELECT * FROM users WHERE username = ?',
            (username,)
        ).fetchone()
        conn.close()
        
        if user and user['password_hash'] == hash_password(password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            return redirect(url_for('index'))
        else:
            return render_template('login.html', error='Неверный логин или пароль')
    
    return render_template('login.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        
        if not username or not password:
            return render_template('register.html', error='Заполните все поля')
        
        if len(password) < 4:
            return render_template('register.html', error='Пароль минимум 4 символа')
        
        conn = get_db()
        try:
            conn.execute(
                'INSERT INTO users (username, password_hash) VALUES (?, ?)',
                (username, hash_password(password))
            )
            conn.commit()
            conn.close()
            return redirect(url_for('login'))
        except sqlite3.IntegrityError:
            conn.close()
            return render_template('register.html', error='Пользователь уже существует')
    
    return render_template('register.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# === API ДЛЯ ДАННЫХ ===

@app.route('/api/day/<date>')
@login_required
def get_day(date):
    """Получить данные за день"""
    conn = get_db()
    record = conn.execute(
        'SELECT * FROM records WHERE user_id = ? AND date = ?',
        (session['user_id'], date)
    ).fetchone()
    conn.close()
    
    if record:
        return jsonify({
            'work': record['work'],
            'hours': record['hours'],
            'salary': record['salary'],
            'salary_day': bool(record['salary_day']),
            'comment': record['comment']
        })
    return jsonify({'work': None, 'hours': None, 'salary': None, 'salary_day': False, 'comment': ''})

@app.route('/api/save', methods=['POST'])
@login_required
def save_day():
    """Сохранить данные за день"""
    data = request.json
    date = data.get('date')
    field = data.get('field')
    value = data.get('value')
    
    if not date or not field:
        return jsonify({'success': False, 'error': 'Нет данных'})
    
    conn = get_db()
    
    # Проверяем, есть ли запись
    record = conn.execute(
        'SELECT id FROM records WHERE user_id = ? AND date = ?',
        (session['user_id'], date)
    ).fetchone()
    
    if record:
        # Обновляем
        if field == 'salary_day':
            conn.execute(
                f'UPDATE records SET {field} = ? WHERE user_id = ? AND date = ?',
                (1 if value else 0, session['user_id'], date)
            )
        else:
            conn.execute(
                f'UPDATE records SET {field} = ? WHERE user_id = ? AND date = ?',
                (value, session['user_id'], date)
            )
    else:
        # Создаём
        conn.execute(
            'INSERT INTO records (user_id, date, ' + field + ') VALUES (?, ?, ?)',
            (session['user_id'], date, value)
        )
    
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/stats/<int:year>/<int:month>')
@login_required
def get_stats(year, month):
    """Получить статистику за месяц"""
    conn = get_db()
    
    days_in_month = calendar.monthrange(year, month)[1]
    work_days = 0
    missed_days = 0
    total_hours = 0
    month_income = 0
    salary_days = 0
    
    for day in range(1, days_in_month + 1):
        date_str = f"{year}-{month:02d}-{day:02d}"
        record = conn.execute(
            'SELECT * FROM records WHERE user_id = ? AND date = ?',
            (session['user_id'], date_str)
        ).fetchone()
        
        if record:
            if record['work'] == 'yes':
                work_days += 1
            elif record['work'] == 'no':
                missed_days += 1
            
            if record['hours']:
                total_hours += record['hours']
            
            if record['salary']:
                month_income += record['salary']
            
            if record['salary_day']:
                salary_days += 1
    
    # Общий доход
    total_income = conn.execute(
        'SELECT SUM(salary) FROM records WHERE user_id = ?',
        (session['user_id'],)
    ).fetchone()[0] or 0
    
    conn.close()
    
    avg_hours = total_hours / work_days if work_days > 0 else 0
    
    return jsonify({
        'days_in_month': days_in_month,
        'work_days': work_days,
        'missed_days': missed_days,
        'total_hours': total_hours,
        'avg_hours': avg_hours,
        'month_income': month_income,
        'total_income': total_income,
        'salary_days': salary_days
    })
@app.route('/api/reset_day', methods=['POST'])
@login_required
def reset_day():
    """Сбросить данные за день"""
    data = request.json
    date = data.get('date')
    
    if not date:
        return jsonify({'success': False, 'error': 'Нет даты'})
    
    conn = get_db()
    conn.execute(
        'DELETE FROM records WHERE user_id = ? AND date = ?',
        (session['user_id'], date)
    )
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/clear_all', methods=['POST'])
@login_required
def clear_all():
    """Удалить все данные пользователя"""
    conn = get_db()
    conn.execute(
        'DELETE FROM records WHERE user_id = ?',
        (session['user_id'],)
    )
    conn.commit()
    conn.close()
    return jsonify({'success': True})

# === ВЫГРУЗКА В EXCEL ===

@app.route('/export/month/<int:year>/<int:month>')
@login_required
def export_month(year, month):
    """Выгрузка за месяц"""
    conn = get_db()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{calendar.month_name[month]} {year}"
    
    headers = ['Дата', 'Статус', 'Часы', 'Зарплата (руб)', 'День зарплаты', 'Комментарий']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
    
    days_in_month = calendar.monthrange(year, month)[1]
    row = 2
    work_days = total_hours = total_salary = salary_days = 0
    
    for day in range(1, days_in_month + 1):
        date_obj = datetime(year, month, day).date()
        date_str = date_obj.strftime("%Y-%m-%d")
        
        record = conn.execute(
            'SELECT * FROM records WHERE user_id = ? AND date = ?',
            (session['user_id'], date_str)
        ).fetchone()
        
        if record:
            status = 'Был' if record['work'] == 'yes' else 'Не был' if record['work'] == 'no' else ''
            hours = record['hours'] or ''
            salary = record['salary'] or ''
            is_salary_day = 'Да' if record['salary_day'] else 'Нет'
            comment = record['comment'] or ''
            
            if record['work'] == 'yes':
                work_days += 1
            if record['hours']:
                total_hours += record['hours']
            if record['salary']:
                total_salary += record['salary']
            if record['salary_day']:
                salary_days += 1
        else:
            status = hours = salary = comment = ''
            is_salary_day = 'Нет'
        
        ws.cell(row=row, column=1, value=date_obj.strftime("%d.%m.%Y"))
        ws.cell(row=row, column=2, value=status)
        ws.cell(row=row, column=3, value=hours)
        ws.cell(row=row, column=4, value=salary)
        ws.cell(row=row, column=5, value=is_salary_day)
        ws.cell(row=row, column=6, value=comment)
        row += 1
    
    # Итого
    row += 1
    ws.cell(row=row, column=1, value="ИТОГО:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=f"{work_days} дн.")
    ws.cell(row=row, column=3, value=f"{total_hours:.1f}")
    ws.cell(row=row, column=4, value=f"{total_salary:,.0f}")
    ws.cell(row=row, column=5, value=f"{salary_days} дн.")
    
    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 18
    
    conn.close()
    
    # Сохраняем в память
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Зарплата_{calendar.month_name[month]}_{year}.xlsx"
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/export/year/<int:year>')
@login_required
def export_year(year):
    """Выгрузка за год"""
    conn = get_db()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}"
    
    headers = ['Месяц', 'Отработано дн.', 'Всего часов', 'Доход (руб)', 'Дней зарплаты']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
    
    total_work_days = total_hours_all = total_salary_all = total_salary_days = 0
    
    for month in range(1, 13):
        days_in_month = calendar.monthrange(year, month)[1]
        work_days = total_hours = total_salary = salary_days = 0
        
        for day in range(1, days_in_month + 1):
            date_str = f"{year}-{month:02d}-{day:02d}"
            record = conn.execute(
                'SELECT * FROM records WHERE user_id = ? AND date = ?',
                (session['user_id'], date_str)
            ).fetchone()
            
            if record:
                if record['work'] == 'yes':
                    work_days += 1
                if record['hours']:
                    total_hours += record['hours']
                if record['salary']:
                    total_salary += record['salary']
                if record['salary_day']:
                    salary_days += 1
        
        ws.cell(row=month + 1, column=1, value=calendar.month_name[month])
        ws.cell(row=month + 1, column=2, value=work_days)
        ws.cell(row=month + 1, column=3, value=f"{total_hours:.1f}")
        ws.cell(row=month + 1, column=4, value=f"{total_salary:,.0f}")
        ws.cell(row=month + 1, column=5, value=salary_days)
        
        total_work_days += work_days
        total_hours_all += total_hours
        total_salary_all += total_salary
        total_salary_days += salary_days
    
    row = 14
    ws.cell(row=row, column=1, value="ИТОГО:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_work_days)
    ws.cell(row=row, column=3, value=f"{total_hours_all:.1f}")
    ws.cell(row=row, column=4, value=f"{total_salary_all:,.0f}")
    ws.cell(row=row, column=5, value=total_salary_days)
    
    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 18
    
    conn.close()
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Зарплата_{year}.xlsx"
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

# === ЗАПУСК ===

if __name__ == '__main__':
    init_db()
    app.run(debug=True, host='0.0.0.0', port=5002)
